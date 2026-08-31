from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any, Iterable
import zipfile

import fitz
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .image_metrics import compare
from .image_similarity_gate import assess_raster_gate


CANVAS_WIDTH = 1600
CANVAS_HEIGHT = 900
MIN_APPROVED_RESOURCE_VISIBLE_FRACTION = 0.5
MIN_APPROVED_RESOURCE_VISIBLE_PIXELS = 8.0
OPC_CONTENT_TYPES_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/content-types"
OPC_RELATIONSHIPS_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"


class PptxError(RuntimeError):
    pass


def resolve_soffice(command: str | Path | None = None) -> Path:
    """Resolve the pinned headless renderer used by the PPTBench benchmark."""
    candidates: list[str | Path] = []
    if command:
        candidates.append(command)
    if os.environ.get("FRONTEND_BENCH_SOFFICE"):
        candidates.append(os.environ["FRONTEND_BENCH_SOFFICE"])
    candidates.extend(["soffice", "libreoffice"])
    for ancestor in Path(__file__).resolve().parents:
        candidates.extend(
            [
                ancestor / ".tools/libreoffice-portable/opt/libreoffice25.8/program/soffice",
                ancestor / ".tools/libreoffice/usr/lib/libreoffice/program/soffice",
            ]
        )
    for candidate in candidates:
        resolved = shutil.which(str(candidate))
        if resolved:
            return Path(resolved).resolve()
        path = Path(candidate).expanduser()
        if path.is_file() and os.access(path, os.X_OK):
            return path.resolve()
    raise PptxError(
        "PPTX rendering requires LibreOffice (soffice). Install it, pass "
        "--soffice-command, or set FRONTEND_BENCH_SOFFICE."
    )


def _normalize_render(source: Path, destination: Path) -> None:
    """Normalize a rendered image to the benchmark canvas without stretching."""
    with Image.open(source) as loaded:
        image = loaded.convert("RGB")
        image.thumbnail((CANVAS_WIDTH, CANVAS_HEIGHT), Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (CANVAS_WIDTH, CANVAS_HEIGHT), "white")
        canvas.paste(
            image,
            ((CANVAS_WIDTH - image.width) // 2, (CANVAS_HEIGHT - image.height) // 2),
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        canvas.save(destination, optimize=True)


def _xml_infoset(element: etree._Element) -> tuple[Any, ...]:
    """Return a prefix-independent representation of a small OPC XML tree."""

    return (
        element.tag,
        tuple(sorted(element.attrib.items())),
        element.text or "",
        element.tail or "",
        tuple(_xml_infoset(child) for child in element),
    )


def _clone_with_default_namespace(
    element: etree._Element,
    namespace: str,
    *,
    root: bool = True,
) -> etree._Element:
    clone = etree.Element(element.tag, nsmap={None: namespace} if root else None)
    for key, value in element.attrib.items():
        clone.set(key, value)
    clone.text = element.text
    clone.tail = element.tail
    for child in element:
        clone.append(_clone_with_default_namespace(child, namespace, root=False))
    return clone


def _write_libreoffice_compatible_opc_copy(
    source: Path,
    destination: Path,
) -> list[str]:
    """Write a render-only copy with default OPC metadata namespaces.

    LibreOffice 25.8 rejects otherwise valid packages when standard content-type or
    relationship elements use a generated namespace prefix. Only those metadata
    members are rewritten, and their expanded XML infosets must remain identical.
    """

    changed: list[str] = []
    expected_roots = {
        "[Content_Types].xml": (
            OPC_CONTENT_TYPES_NAMESPACE,
            f"{{{OPC_CONTENT_TYPES_NAMESPACE}}}Types",
        )
    }
    try:
        with zipfile.ZipFile(source) as archive:
            infos = archive.infolist()
            names = [info.filename for info in infos]
            if len(names) != len(set(names)):
                raise PptxError("PPTX package contains duplicate ZIP members")
            destination.parent.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(destination, "w") as output:
                output.comment = archive.comment
                for info in infos:
                    data = archive.read(info)
                    expected = expected_roots.get(info.filename)
                    if expected is None and info.filename.endswith(".rels"):
                        expected = (
                            OPC_RELATIONSHIPS_NAMESPACE,
                            f"{{{OPC_RELATIONSHIPS_NAMESPACE}}}Relationships",
                        )
                    if expected is not None:
                        namespace, root_tag = expected
                        parsed = etree.fromstring(data)
                        prefixed = any(
                            child.prefix is not None
                            for child in parsed.iter()
                            if isinstance(child.tag, str)
                        )
                        if parsed.tag == root_tag and prefixed:
                            normalized_root = _clone_with_default_namespace(parsed, namespace)
                            normalized = etree.tostring(
                                normalized_root,
                                encoding="UTF-8",
                                xml_declaration=True,
                            )
                            reparsed = etree.fromstring(normalized)
                            if _xml_infoset(parsed) != _xml_infoset(reparsed):
                                raise PptxError(
                                    "OPC namespace normalization changed the XML infoset"
                                )
                            data = normalized
                            changed.append(info.filename)
                    output.writestr(info, data)
    except (OSError, zipfile.BadZipFile, etree.XMLSyntaxError) as exc:
        destination.unlink(missing_ok=True)
        raise PptxError(f"cannot normalize OPC metadata for rendering: {exc}") from exc

    if not changed:
        destination.unlink(missing_ok=True)
        return []

    try:
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination) as copy:
            original_names = [info.filename for info in original.infolist()]
            copy_names = [info.filename for info in copy.infolist()]
            if copy_names != original_names:
                raise PptxError("render-only OPC copy changed the ZIP member order")
            changed_set = set(changed)
            for name in original_names:
                before = original.read(name)
                after = copy.read(name)
                if name not in changed_set:
                    if before != after:
                        raise PptxError(f"render-only OPC copy changed unexpected member: {name}")
                    continue
                if _xml_infoset(etree.fromstring(before)) != _xml_infoset(etree.fromstring(after)):
                    raise PptxError(f"render-only OPC copy changed metadata semantics: {name}")
    except (OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
        destination.unlink(missing_ok=True)
        raise
    return changed


def _run_soffice_render_attempt(
    *,
    soffice: Path,
    pptx_path: Path,
    root: Path,
    label: str,
    render_env: dict[str, str],
) -> tuple[list[str], subprocess.CompletedProcess[str], Path]:
    out = root / f"{label}-out"
    profile = root / f"{label}-profile"
    out.mkdir()
    profile.mkdir()
    command = [
        str(soffice),
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(out),
        str(pptx_path.resolve()),
    ]
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        check=False,
        timeout=120,
        env=render_env,
    )
    return command, result, out / f"{pptx_path.stem}.pdf"


def render_pptx(
    pptx_path: Path,
    screenshot_path: Path,
    *,
    soffice_command: str | Path | None = None,
) -> None:
    """Render the first PPTX slide through LibreOffice and normalize it to 1600x900."""
    soffice = resolve_soffice(soffice_command)
    screenshot_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pptbench-pptx-") as temporary:
        root = Path(temporary)
        local_library_dirs = [soffice.parent]
        for ancestor in Path(__file__).resolve().parents:
            local_library_dirs.append(ancestor / ".tools/libreoffice/usr/lib/x86_64-linux-gnu")
        library_path = [str(path) for path in local_library_dirs if path.is_dir()]
        if os.environ.get("LD_LIBRARY_PATH"):
            library_path.append(os.environ["LD_LIBRARY_PATH"])
        render_env = {
            **os.environ,
            **({"LD_LIBRARY_PATH": os.pathsep.join(library_path)} if library_path else {}),
        }
        command, result, rendered_pdf = _run_soffice_render_attempt(
            soffice=soffice,
            pptx_path=pptx_path,
            root=root,
            label="attempt-1",
            render_env=render_env,
        )
        attempts = ["$ " + " ".join(command) + "\n" + result.stdout + result.stderr]
        combined_output = result.stdout + result.stderr
        if (
            result.returncode == 0
            and not rendered_pdf.is_file()
            and "Error: source file could not be loaded" in combined_output
        ):
            compatible_copy = root / "render-compatible.pptx"
            try:
                changed_members = _write_libreoffice_compatible_opc_copy(pptx_path, compatible_copy)
            except PptxError as exc:
                attempts.append(f"OPC fallback rejected: {exc}\n")
                changed_members = []
            if changed_members:
                attempts.append(
                    "OPC fallback source_sha256="
                    + _sha256_bytes(pptx_path.read_bytes())
                    + " copy_sha256="
                    + _sha256_bytes(compatible_copy.read_bytes())
                    + " changed_members="
                    + json.dumps(changed_members)
                    + "\n"
                )
                command, result, rendered_pdf = _run_soffice_render_attempt(
                    soffice=soffice,
                    pptx_path=compatible_copy,
                    root=root,
                    label="attempt-2",
                    render_env=render_env,
                )
                attempts.append("$ " + " ".join(command) + "\n" + result.stdout + result.stderr)
                combined_output = result.stdout + result.stderr
        (screenshot_path.parent / "render.log").write_text(
            "\n".join(attempts),
            encoding="utf-8",
        )
        if result.returncode or not rendered_pdf.is_file():
            raise PptxError(
                f"LibreOffice PPTX rendering failed with code {result.returncode}: "
                f"{combined_output[-1200:]}"
            )
        with fitz.open(rendered_pdf) as document:
            if len(document) != 1:
                raise PptxError(
                    f"reconstruction.pptx must render to exactly one slide, got {len(document)}"
                )
            page = document[0]
            scale = max(2.0, 2000 / max(page.rect.width, page.rect.height))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            raw = root / "slide.png"
            pixmap.save(raw)
        _normalize_render(raw, screenshot_path)


def _color(value: Any) -> str | None:
    """Extract a normalized hexadecimal color from a shape property."""
    try:
        if value is None or value.type is None:
            return None
        if value.type.name == "RGB":
            return f"#{value.rgb}"
        if value.type.name == "THEME":
            return f"theme:{value.theme_color}"
        return str(value.type)
    except (AttributeError, TypeError, ValueError):
        return None


def _shape_kind(shape: Any) -> str:
    """Return a stable shape-kind label."""
    value = shape.shape_type
    return getattr(value, "name", str(value)).lower()


def _shape_text(shape: Any) -> str:
    """Extract visible text from a shape."""
    if not getattr(shape, "has_text_frame", False):
        return ""
    paragraphs = [paragraph.text.strip() for paragraph in shape.text_frame.paragraphs]
    return "\n".join(value for value in paragraphs if value)


def _shape_style(shape: Any) -> dict[str, Any]:
    """Extract fill, line, and text styling from a shape."""
    style: dict[str, Any] = {}
    try:
        style["fill"] = _color(shape.fill.fore_color) if shape.fill.type is not None else None
    except (AttributeError, TypeError, ValueError):
        style["fill"] = None
    try:
        style["line"] = _color(shape.line.color)
        style["line_width_pt"] = (
            round(float(shape.line.width.pt), 3) if shape.line.width is not None else None
        )
    except (AttributeError, TypeError, ValueError):
        style["line"] = None
        style["line_width_pt"] = None
    return style


def _flatten_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    """Yield nested group shapes recursively."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _flatten_shapes(shape.shapes)
        else:
            yield shape


def _canvas_bbox(
    left: float,
    top: float,
    width: float,
    height: float,
    source_width: float,
    source_height: float,
) -> list[float]:
    """Map a shape bounding box into normalized canvas coordinates."""
    scale = min(CANVAS_WIDTH / source_width, CANVAS_HEIGHT / source_height)
    offset_x = (CANVAS_WIDTH - source_width * scale) / 2
    offset_y = (CANVAS_HEIGHT - source_height * scale) / 2
    return [
        round(offset_x + left * scale, 3),
        round(offset_y + top * scale, 3),
        round(offset_x + (left + width) * scale, 3),
        round(offset_y + (top + height) * scale, 3),
    ]


def _external_relationships(pptx_path: Path) -> list[str]:
    """Return external relationships declared by a PPTX package."""
    external: list[str] = []
    with zipfile.ZipFile(pptx_path) as archive:
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            text = archive.read(name).decode("utf-8", errors="replace")
            external.extend(
                re.findall(
                    r'<Relationship\b[^>]*\bTargetMode="External"[^>]*\bTarget="([^"]+)"', text
                )
            )
            external.extend(
                re.findall(
                    r'<Relationship\b[^>]*\bTarget="([^"]+)"[^>]*\bTargetMode="External"', text
                )
            )
    return sorted(set(external))


def _sha256_bytes(value: bytes) -> str:
    """Return the SHA-256 digest of an in-memory package part."""

    return hashlib.sha256(value).hexdigest()


def _shape_is_hidden(shape: Any) -> bool:
    """Return whether OOXML explicitly marks a shape as hidden."""

    try:
        properties = shape.element.xpath(".//p:cNvPr")
    except (AttributeError, KeyError):
        return False
    if not properties:
        return False
    return str(properties[0].get("hidden", "")).casefold() in {"1", "true"}


def _picture_visibility(bbox: list[float]) -> dict[str, float]:
    """Measure how much of a picture's bounding box intersects the slide."""

    left, top, right, bottom = bbox
    width = max(0.0, right - left)
    height = max(0.0, bottom - top)
    visible_width = max(0.0, min(1.0, right) - max(0.0, left))
    visible_height = max(0.0, min(1.0, bottom) - max(0.0, top))
    area = width * height
    visible_area = visible_width * visible_height
    return {
        "visible_intersection_ratio": round(visible_area / area, 6) if area else 0.0,
        "visible_area_coverage": round(visible_area, 6),
        "visible_width_pixels": round(visible_width * CANVAS_WIDTH, 3),
        "visible_height_pixels": round(visible_height * CANVAS_HEIGHT, 3),
    }


def _extract_pptx_components(pptx_path: Path) -> dict[str, Any]:
    """Extract components after the public wrapper has established error semantics."""
    if not pptx_path.is_file() or not zipfile.is_zipfile(pptx_path):
        raise PptxError("reconstruction.pptx is missing or is not a valid OOXML zip")
    try:
        presentation = Presentation(pptx_path)
    except Exception as exc:  # python-pptx raises several package/XML exceptions
        raise PptxError(f"cannot parse reconstruction.pptx: {exc}") from exc
    slide_width = float(presentation.slide_width)
    slide_height = float(presentation.slide_height)
    components: list[dict[str, Any]] = []
    picture_count = 0
    slide_image_references: list[dict[str, Any]] = []
    editable_count = 0
    connector_count = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for relationship in slide.part.rels.values():
            if relationship.reltype != RT.IMAGE:
                continue
            if relationship.is_external:
                slide_image_references.append(
                    {
                        "slide": slide_index,
                        "relationship_id": relationship.rId,
                        "external": True,
                        "target": relationship.target_ref,
                    }
                )
                continue
            blob = relationship.target_part.blob
            slide_image_references.append(
                {
                    "slide": slide_index,
                    "relationship_id": relationship.rId,
                    "content_type": relationship.target_part.content_type,
                    "image_bytes": len(blob),
                    "image_sha256": _sha256_bytes(blob),
                }
            )
        for z_index, shape in enumerate(_flatten_shapes(slide.shapes)):
            left = float(shape.left)
            top = float(shape.top)
            width = float(shape.width)
            height = float(shape.height)
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            is_connector = shape.shape_type == MSO_SHAPE_TYPE.LINE or shape.element.tag.endswith(
                "}cxnSp"
            )
            if is_picture:
                picture_count += 1
            else:
                editable_count += 1
            if is_connector:
                connector_count += 1
            component: dict[str, Any] = {
                "component_id": f"s{slide_index}-o{z_index + 1}",
                "slide": slide_index,
                "z_index": z_index,
                "name": shape.name,
                "kind": _shape_kind(shape),
                "editable": not is_picture,
                "connector": is_connector,
                "bbox": _canvas_bbox(left, top, width, height, slide_width, slide_height),
                "bbox_normalized": [
                    round(left / slide_width, 6),
                    round(top / slide_height, 6),
                    round((left + width) / slide_width, 6),
                    round((top + height) / slide_height, 6),
                ],
                "rotation": float(shape.rotation or 0),
                "text": _shape_text(shape),
                "style": _shape_style(shape),
            }
            if is_picture:
                component["hidden"] = _shape_is_hidden(shape)
                component.update(_picture_visibility(component["bbox_normalized"]))
                try:
                    image_blob = shape.image.blob
                    component["image_bytes"] = len(image_blob)
                    component["image_sha1"] = shape.image.sha1
                    component["image_sha256"] = _sha256_bytes(image_blob)
                except (AttributeError, KeyError):
                    pass
            components.append(component)
    total_count = len(components)
    report = {
        "schema": "pptbench-pptx-components-v1",
        "source_path": str(pptx_path.resolve()),
        "slide_count": len(presentation.slides),
        "slide_size": {
            "width_emu": int(slide_width),
            "height_emu": int(slide_height),
            "aspect_ratio": round(slide_width / slide_height, 6),
        },
        "summary": {
            "component_count": total_count,
            "editable_component_count": editable_count,
            "picture_count": picture_count,
            "connector_count": connector_count,
            "text_component_count": sum(bool(row["text"]) for row in components),
            "editability_ratio": round(editable_count / max(1, total_count), 6),
            "external_relationships": _external_relationships(pptx_path),
        },
        "slide_image_references": slide_image_references,
        "components": components,
    }
    return report


def _canonicalize_integral_emu_spellings(pptx_path: Path) -> int:
    """Rewrite only numerically integral OOXML coordinates to integer text.

    Some authoring code serializes an exact EMU value as ``"2880360.0"``.
    Office applications treat that as the same coordinate, but python-pptx's
    lazy simple-type parser calls ``int(value)`` and rejects the spelling.  A
    non-integral value such as ``"1.24996875"`` is not repaired because it is
    not a valid EMU coordinate and remains a malformed model artifact.
    """

    if not pptx_path.is_file() or not zipfile.is_zipfile(pptx_path):
        return 0
    with zipfile.ZipFile(pptx_path, "r") as source:
        members = [(item, source.read(item.filename)) for item in source.infolist()]

    coordinate_attributes = {
        "off": ("x", "y"),
        "ext": ("cx", "cy"),
        "chOff": ("x", "y"),
        "chExt": ("cx", "cy"),
        "sldSz": ("cx", "cy"),
        "notesSz": ("cx", "cy"),
    }
    replacements: dict[str, bytes] = {}
    rewritten = 0
    for item, payload in members:
        if not item.filename.endswith(".xml"):
            continue
        try:
            root = etree.fromstring(payload)
        except etree.XMLSyntaxError:
            continue
        changed = False
        for element in root.iter():
            local_name = etree.QName(element).localname
            for attribute in coordinate_attributes.get(local_name, ()):
                raw = element.get(attribute)
                if raw is None or "." not in raw:
                    continue
                try:
                    value = Decimal(raw)
                except InvalidOperation:
                    continue
                if not value.is_finite() or value != value.to_integral_value():
                    continue
                canonical = str(int(value))
                if canonical == raw:
                    continue
                element.set(attribute, canonical)
                rewritten += 1
                changed = True
        if changed:
            replacements[item.filename] = etree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=payload.lstrip().startswith(b"<?xml"),
            )

    if not replacements:
        return 0
    with tempfile.NamedTemporaryFile(
        prefix=f".{pptx_path.stem}-emu-", suffix=".pptx", dir=pptx_path.parent, delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w") as destination:
            for item, payload in members:
                destination.writestr(item, replacements.get(item.filename, payload))
        temporary.replace(pptx_path)
    finally:
        temporary.unlink(missing_ok=True)
    return rewritten


def extract_pptx_components(pptx_path: Path) -> dict[str, Any]:
    """Extract editable objects while normalizing lazy OOXML parse failures.

    ``python-pptx`` loads portions of a deck lazily, so malformed coordinates
    and style values may not fail during ``Presentation(...)`` construction.
    They can instead raise ``ValueError`` (or another package exception) when a
    shape property is first accessed.  Convert all such failures to
    :class:`PptxError` so the matrix records a malformed model artifact rather
    than retrying it as an evaluator outage.
    """

    try:
        return _extract_pptx_components(pptx_path)
    except Exception as first_error:  # python-pptx lazily parses shape XML
        try:
            rewritten = _canonicalize_integral_emu_spellings(pptx_path)
        except Exception:
            rewritten = 0
        if rewritten:
            try:
                return _extract_pptx_components(pptx_path)
            except Exception as retry_error:
                if isinstance(retry_error, PptxError):
                    raise
                raise PptxError(
                    f"cannot parse reconstruction.pptx: {retry_error}"
                ) from retry_error
        if isinstance(first_error, PptxError):
            raise
        raise PptxError(f"cannot parse reconstruction.pptx: {first_error}") from first_error


def audit_pptx_resource_usage(
    report: dict[str, Any],
    approved_resource_paths: Iterable[Path],
) -> dict[str, Any]:
    """Check approved resource usage and external image relationships.

    Approved source assets are required to be embedded byte-for-byte and used
    visibly.  Other image media is retained for the image-similarity artifact
    gate; it is not rejected merely because it is raster.  This separation lets
    the similarity gate distinguish a legitimate small image from a copied
    reference while keeping resource-contract diagnostics useful.
    """

    approved_by_sha: dict[str, list[str]] = {}
    approved_records: list[dict[str, Any]] = []
    for path in sorted((Path(value) for value in approved_resource_paths), key=str):
        if not path.is_file():
            raise PptxError(f"approved raster resource is missing: {path}")
        blob = path.read_bytes()
        digest = _sha256_bytes(blob)
        approved_by_sha.setdefault(digest, []).append(path.name)
        approved_records.append(
            {
                "file": path.name,
                "bytes": len(blob),
                "sha256": digest,
            }
        )

    references = report.get("slide_image_references", [])
    referenced_hashes = {
        str(row.get("image_sha256", ""))
        for row in references
        if isinstance(row, dict) and row.get("image_sha256")
    }
    approved_hashes = set(approved_by_sha)
    violations: list[str] = []
    external_references = [
        row for row in references if isinstance(row, dict) and row.get("external")
    ]
    violations.extend(
        "slide references external image media: " + str(row.get("target", ""))
        for row in external_references
    )
    picture_uses = [
        component
        for component in report.get("components", [])
        if isinstance(component, dict) and component.get("kind") == "picture"
    ]
    meaningful_hashes = {
        str(component.get("image_sha256"))
        for component in picture_uses
        if component.get("image_sha256")
        and not component.get("hidden", False)
        and float(component.get("visible_intersection_ratio", 0.0))
        >= MIN_APPROVED_RESOURCE_VISIBLE_FRACTION
        and float(component.get("visible_width_pixels", 0.0))
        >= MIN_APPROVED_RESOURCE_VISIBLE_PIXELS
        and float(component.get("visible_height_pixels", 0.0))
        >= MIN_APPROVED_RESOURCE_VISIBLE_PIXELS
    }
    for digest in sorted(approved_hashes - meaningful_hashes):
        names = ", ".join(sorted(approved_by_sha[digest]))
        violations.append(
            f"approved raster resource is not visibly used on-slide: {names} (SHA-256 {digest})"
        )

    return {
        "schema": "pptbench-pptx-resource-integrity-v1",
        "approved_resources": approved_records,
        "approved_resource_count": len(approved_records),
        "approved_unique_sha256_count": len(approved_hashes),
        "slide_image_reference_count": len(references),
        "referenced_sha256": sorted(referenced_hashes),
        "meaningfully_visible_sha256": sorted(meaningful_hashes),
        "minimum_visible_intersection_ratio": MIN_APPROVED_RESOURCE_VISIBLE_FRACTION,
        "minimum_visible_width_pixels": MIN_APPROVED_RESOURCE_VISIBLE_PIXELS,
        "minimum_visible_height_pixels": MIN_APPROVED_RESOURCE_VISIBLE_PIXELS,
        "violations": violations,
        "status": "complete" if not violations else "invalid",
    }


def validate_pptx_report(report: dict[str, Any]) -> list[str]:
    """Validate the structural artifact report and return violations."""
    summary = report["summary"]
    violations: list[str] = []
    if report["slide_count"] != 1:
        violations.append(f"deck must contain exactly one slide (got {report['slide_count']})")
    # A one-slide package with no meaningful editable content is not a
    # reconstruction.  This is an artifact-contract sanity check; anti-copy
    # decisions are made solely by the per-image similarity gate below, not by
    # a native-object count or a slide-area heuristic.
    if summary["editable_component_count"] < 3:
        violations.append("deck has fewer than three editable components")
    if summary["external_relationships"]:
        violations.append("deck contains external relationships")
    return violations


def validate_reconstruction_pptx(
    pptx_path: Path,
    *,
    reference_path: Path | None = None,
    approved_resource_paths: Iterable[Path] | None = None,
    enforce_resource_integrity: bool = True,
    enforce_image_similarity: bool = True,
) -> dict[str, Any]:
    """Validate a candidate reconstruction PPTX before scoring."""
    try:
        with zipfile.ZipFile(pptx_path) as archive:
            forbidden_payloads = [
                name
                for name in archive.namelist()
                if name.startswith(("ppt/fonts/", "ppt/embeddings/", "ppt/activeX/"))
            ]
    except (OSError, zipfile.BadZipFile) as exc:
        raise PptxError(f"invalid PPTX package: {exc}") from exc
    if forbidden_payloads:
        raise PptxError(
            "invalid direct PPTX reconstruction: deck contains embedded binary payloads: "
            + ", ".join(forbidden_payloads)
        )
    report = extract_pptx_components(pptx_path)
    violations = validate_pptx_report(report)
    approved_hashes: list[str] = []
    if approved_resource_paths is not None:
        resource_integrity = audit_pptx_resource_usage(report, approved_resource_paths)
        report["resource_integrity"] = resource_integrity
        approved_hashes = [
            str(row.get("sha256", ""))
            for row in resource_integrity.get("approved_resources", [])
            if row.get("sha256")
        ]
        if enforce_resource_integrity:
            violations.extend(resource_integrity["violations"])
    if reference_path is not None:
        if not Path(reference_path).is_file():
            raise PptxError(f"reference image is missing: {reference_path}")
        similarity_gate = assess_raster_gate(
            Path(reference_path),
            None,
            report,
            approved_resource_hashes=approved_hashes,
        )
        report["image_similarity_gate"] = similarity_gate
        if enforce_image_similarity:
            violations.extend(similarity_gate["violations"])
    else:
        report["image_similarity_gate"] = {
            "status": "not_run",
            "reason": "reference image was not supplied",
            "triggered": False,
            "violations": [],
        }
    if violations:
        raise PptxError("invalid direct PPTX reconstruction: " + "; ".join(violations))
    return report


def _tokens(value: str) -> list[str]:
    """Tokenize text for content matching."""
    return re.findall(r"[\w]+", value.casefold(), flags=re.UNICODE)


def _counter_f1(reference: Counter[str], candidate: Counter[str]) -> float:
    """Compute token-count F1 between reference and candidate text."""
    overlap = sum((reference & candidate).values())
    precision = overlap / max(1, sum(candidate.values()))
    recall = overlap / max(1, sum(reference.values()))
    return 2 * precision * recall / max(1e-9, precision + recall)


def _bbox_center(bbox: list[float]) -> tuple[float, float]:
    """Return the center of a normalized bounding box."""
    return ((bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2)


def _pdf_text_components(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    """Read reference text components from the structural manifest."""
    if not manifest.get("pages"):
        return []
    page = manifest["pages"][0]
    width = float(page["width"])
    height = float(page["height"])
    result = []
    for item in page.get("texts", []):
        text = str(item.get("text", "")).strip()
        if not _tokens(text):
            continue
        x1, y1, x2, y2 = (float(value) for value in item["bbox"])
        result.append(
            {
                "component_id": item.get("component_id"),
                "text": text,
                "bbox": _canvas_bbox(x1, y1, x2 - x1, y2 - y1, width, height),
            }
        )
    return result


def inspect_reference_pdf(source_pdf_path: Path) -> dict[str, Any]:
    """Build the structural reference needed by PPTX scoring from page one."""

    try:
        document = fitz.open(source_pdf_path)
    except (OSError, RuntimeError, ValueError) as exc:
        raise PptxError(f"cannot inspect source PDF {source_pdf_path}: {exc}") from exc
    try:
        if document.page_count < 1:
            raise PptxError(f"source PDF has no pages: {source_pdf_path}")
        page = document[0]
        text_dict = page.get_text("dict")
        texts: list[dict[str, Any]] = []
        for block_index, block in enumerate(text_dict.get("blocks", [])):
            if block.get("type") != 0:
                continue
            for line_index, line in enumerate(block.get("lines", [])):
                for span_index, span in enumerate(line.get("spans", [])):
                    text = str(span.get("text", "")).strip()
                    bbox = span.get("bbox")
                    if not text or not isinstance(bbox, (list, tuple)) or len(bbox) != 4:
                        continue
                    texts.append(
                        {
                            "component_id": (
                                f"p1-b{block_index + 1}-l{line_index + 1}-s{span_index + 1}"
                            ),
                            "text": text,
                            "bbox": [round(float(value), 4) for value in bbox],
                        }
                    )
        drawings = page.get_drawings()
        connector_count = sum(
            1
            for drawing in drawings
            if any(item and item[0] in {"l", "c", "qu"} for item in drawing.get("items", []))
        )
        return {
            "schema": "pptbench-runtime-pdf-components-v1",
            "source_path": str(source_pdf_path.resolve()),
            "summary": {"connector_count": connector_count},
            "pages": [
                {
                    "page_number": 1,
                    "width": round(float(page.rect.width), 4),
                    "height": round(float(page.rect.height), 4),
                    "texts": texts,
                }
            ],
        }
    finally:
        document.close()


def _positioned_text_score(
    reference: list[dict[str, Any]], candidate: list[dict[str, Any]]
) -> float:
    """Compare text content and positions between reference and candidate."""
    if not reference:
        return 1.0 if not candidate else 0.0
    available = set(range(len(candidate)))
    scores: list[float] = []
    diagonal = math.hypot(CANVAS_WIDTH, CANVAS_HEIGHT)
    for expected in sorted(reference, key=lambda row: len(_tokens(row["text"])), reverse=True):
        best: tuple[float, int] | None = None
        for index in available:
            actual = candidate[index]
            similarity = SequenceMatcher(
                None,
                " ".join(_tokens(expected["text"])),
                " ".join(_tokens(actual["text"])),
            ).ratio()
            if similarity < 0.25:
                continue
            ex, ey = _bbox_center(expected["bbox"])
            ax, ay = _bbox_center(actual["bbox"])
            proximity = max(0.0, 1 - math.hypot(ex - ax, ey - ay) / (0.35 * diagonal))
            score = 0.72 * similarity + 0.28 * proximity
            if best is None or score > best[0]:
                best = (score, index)
        if best is None:
            scores.append(0.0)
        else:
            scores.append(best[0])
            available.remove(best[1])
    return sum(scores) / len(scores)


def _count_similarity(reference: int, candidate: int) -> float:
    """Compare reference and candidate component counts."""
    if reference == candidate:
        return 1.0
    return math.exp(-abs(math.log1p(reference) - math.log1p(candidate)))


def evaluate_pptx(
    reference_path: Path,
    screenshot_path: Path,
    pptx_report: dict[str, Any],
    pdf_component_manifest: Path | None = None,
    source_pdf_path: Path | None = None,
) -> dict[str, Any]:
    """Score a direct PPTX output visually and structurally."""
    visual = compare(reference_path, screenshot_path)
    summary = pptx_report["summary"]
    editable = min(1.0, float(summary["editability_ratio"]))
    # Raster integrity is an artifact gate, not a continuous editability
    # discount.  A near-copy is rejected by ``image_similarity_gate``; a
    # legitimate embedded asset must not lower the structural score merely
    # because it occupies a large rectangle.
    editability_score = editable
    component: dict[str, Any] = {
        "available": False,
        "text_content_score": None,
        "positioned_text_score": None,
        "score": None,
    }
    topology: dict[str, Any] = {
        "available": False,
        "reference_connector_count": None,
        "candidate_connector_count": summary["connector_count"],
        "score": None,
    }
    manifest: dict[str, Any] | None = None
    reference_source = ""
    if pdf_component_manifest is not None and pdf_component_manifest.is_file():
        manifest = json.loads(pdf_component_manifest.read_text(encoding="utf-8"))
        reference_source = "component-manifest"
    elif source_pdf_path is not None and source_pdf_path.is_file():
        manifest = inspect_reference_pdf(source_pdf_path)
        reference_source = "source-pdf"
    if manifest is not None:
        reference_text = _pdf_text_components(manifest)
        candidate_text = [row for row in pptx_report["components"] if row["text"]]
        text_content = _counter_f1(
            Counter(token for row in reference_text for token in _tokens(row["text"])),
            Counter(token for row in candidate_text for token in _tokens(row["text"])),
        )
        positioned_text = _positioned_text_score(reference_text, candidate_text)
        component_score = 0.65 * text_content + 0.35 * positioned_text
        reference_connectors = int(manifest.get("summary", {}).get("connector_count", 0))
        topology_score = _count_similarity(reference_connectors, int(summary["connector_count"]))
        component = {
            "available": True,
            "reference_source": reference_source,
            "reference_text_count": len(reference_text),
            "candidate_text_count": len(candidate_text),
            "text_content_score": round(text_content, 6),
            "positioned_text_score": round(positioned_text, 6),
            "score": round(component_score, 6),
        }
        topology = {
            "available": True,
            "reference_source": reference_source,
            "reference_connector_count": reference_connectors,
            "candidate_connector_count": summary["connector_count"],
            "score": round(topology_score, 6),
            "note": "v1 compares explicit PPT connectors with PDF connector primitives; raw counts are retained for audit.",
        }
    component_score = component["score"] if component["score"] is not None else 0.0
    topology_score = topology["score"] if topology["score"] is not None else 0.0
    benchmark_score = (
        0.60 * float(visual["composite"])
        + 0.20 * editability_score
        + 0.15 * float(component_score)
        + 0.05 * float(topology_score)
    )
    return {
        **visual,
        "visual_score": visual["composite"],
        "editability_score": round(editability_score, 6),
        "component_score": component["score"],
        "topology_score": topology["score"],
        "pptx_benchmark_score": round(benchmark_score, 6),
        "composite": round(benchmark_score, 6),
        "visual": visual,
        "editability": {
            "score": round(editability_score, 6),
            **summary,
        },
        "component": component,
        "topology": topology,
    }
